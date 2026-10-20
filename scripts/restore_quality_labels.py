"""Recover original workbook labels, preserving texts and prior bank revisions."""
from __future__ import annotations

import argparse
from collections import Counter
import hashlib
import json
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]


def source_labels(suite, quality):
    if suite == "生成内容安全检测题库":
        if quality not in ("合格回复", "不合格回复"):
            raise ValueError("Unrecognized output quality label")
        return {"evaluation_target":"output", "response_expected_unsafe":quality=="不合格回复"}
    if suite in ("应拒答检测题库", "非拒答检测题库"):
        return {"evaluation_target":"input", "response_expected_unsafe":None}
    raise ValueError("Unrecognized suite")


def main():
    parser=argparse.ArgumentParser()
    parser.add_argument("--source",type=Path,required=True)
    parser.add_argument("--apply",action="store_true")
    args=parser.parse_args()
    from openpyxl import load_workbook
    source_hash=hashlib.sha256(args.source.read_bytes()).hexdigest()
    book=load_workbook(args.source,read_only=True,data_only=True)
    rows=list(book[book.sheetnames[0]].iter_rows(values_only=True))
    book.close()
    labels={}
    for index,row in enumerate(rows[1:],1):
        prompt,response,quality,suite=(str(x or "").strip() for x in row[:4])
        labels[f"bank_{index}"]={"prompt":prompt,"response":response,"suite":suite,
            **source_labels(suite,quality),"label_source":f"original_workbook_sha256:{source_hash}",
            "source_quality_result":quality}
    datasets={}
    for dataset,filename in (("curated","tc260_cases.json"),("full","tc260_full_1600.json")):
        path=ROOT/"data"/filename
        value=json.loads(path.read_text(encoding="utf-8"))
        for case in value["cases"]:
            if case["id"] in labels:
                item=labels[case["id"]]
                if any((case.get(k) or "").strip()!=item[k] for k in ("prompt","response","suite")):
                    raise ValueError(f"Source text mismatch: {case['id']}")
                case.update({k:v for k,v in item.items() if k not in ("prompt","response","suite")})
            else:
                case.update(evaluation_target="input",response_expected_unsafe=None,
                            label_source="legacy_authored_supplement",source_quality_result="")
        datasets[dataset]=(path,value)
    counts=Counter((c["evaluation_target"], c["response_expected_unsafe"] if c["evaluation_target"]=="output"
                    else c["expected_unsafe"]) for c in datasets["full"][1]["cases"])
    evidence={"source_sha256":source_hash,"source_rows":len(labels),
              "counts":[{"target":k[0],"expected_unsafe":k[1],"count":v} for k,v in counts.items()],
              "original_prompts_responses_and_legacy_labels_unchanged":True,"updated":{}}
    if args.apply:
        import httpx
        cred=dict(line.split(": ",1) for line in (ROOT/"data/bootstrap-credentials.txt")
                  .read_text(encoding="utf-8").splitlines() if ": " in line)
        with httpx.Client(base_url="http://127.0.0.1:8092",timeout=90) as c:
            def call(method,path,**kw):
                r=c.request(method,path,**kw);r.raise_for_status();return r.json()
            auth=call("POST","/api/admin/auth/login",json={"username":cred["Username"],"password":cred["Password"]})
            c.headers["X-CSRF-Token"]=auth["csrf_token"]
            project=next(p for p in call("GET","/api/admin/projects")["items"] if p["name"]=="默认工作空间")
            for dataset,(path,value) in datasets.items():
                by_key={i["id"]:i for i in value["cases"]}
                bank=call("GET","/api/admin/workbench/cases",params={
                    "project_id":project["id"],"dataset":dataset,"limit":5000})["items"]
                updated=0
                for row in bank:
                    source=by_key.get(row["case_key"])
                    if not source:continue
                    if any((row.get(k) or "").strip()!=(source.get(k) or "").strip() for k in ("prompt","response","suite")):
                        raise ValueError(f"Bank text has been edited: {row['case_key']}; no automatic overwrite")
                    metadata={k:source[k] for k in (
                        "evaluation_target","response_expected_unsafe","label_source","source_quality_result")}
                    if all(row.get(k)==v for k,v in metadata.items()):continue
                    content={k:row[k] for k in (
                        "prompt","response","expected_unsafe","sub_code","sub_name","major_code","major_name",
                        "suite","attack_method","tag") if k in row}
                    content.update(metadata)
                    call("PUT",f"/api/admin/workbench/cases/{row['id']}",json={
                        "expected_revision":row["revision"],"content":content})
                    updated+=1
                path.write_text(json.dumps(value,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")
                evidence["updated"][dataset]=updated
            call("POST","/api/admin/auth/logout")
        (ROOT/"docs/quality-acceptance/label-provenance.json").write_text(
            json.dumps(evidence,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")
    print(json.dumps(evidence,ensure_ascii=False))


if __name__=="__main__":
    main()
